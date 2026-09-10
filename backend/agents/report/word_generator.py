from pathlib import Path

from docx import Document
from docx.shared import Inches, Pt


def generate_word(report: dict, output_path: str | Path) -> dict:
    """
    report must be a normalized ReportSpec dictionary.
    """
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Calibri"
    normal_style.font.size = Pt(10)

    document.add_heading(report["title"], level=0)

    if report["subtitle"]:
        document.add_paragraph(report["subtitle"])

    document.add_paragraph(
        f"Prepared by: {report['prepared_by']}\n"
        f"Generated at: {report['generated_at']}"
    )

    if report["approval_note"]:
        document.add_heading("Approval Note", level=1)
        document.add_paragraph("Status: PENDING APPROVAL")
        document.add_paragraph(
            "This document records a request for review. "
            "It does not constitute approval or authorization."
        )
        document.add_paragraph(
            f"Decision requested: {report['decision_requested']}"
        )
        document.add_paragraph("Approver: ____________________")
        document.add_paragraph("Decision: ____________________")
        document.add_paragraph("Date: ________________________")

    if report["summary"]:
        document.add_heading("Executive Summary", level=1)
        document.add_paragraph(report["summary"])

    for section in report["sections"]:
        document.add_heading(section["heading"], level=1)

        for paragraph in section["paragraphs"]:
            document.add_paragraph(paragraph)

        for table_spec in section["tables"]:
            document.add_heading(table_spec["title"], level=2)

            table = document.add_table(
                rows=1,
                cols=len(table_spec["columns"]),
            )
            table.style = "Table Grid"

            for cell, value in zip(
                table.rows[0].cells,
                table_spec["columns"],
            ):
                cell.text = value

            for row in table_spec["rows"]:
                cells = table.add_row().cells

                for cell, value in zip(cells, row):
                    cell.text = value

            if table_spec["note"]:
                document.add_paragraph(table_spec["note"])

        for figure in section["figures"]:
            document.add_picture(
                figure["path"],
                width=Inches(6),
            )
            if figure["caption"]:
                document.add_paragraph(figure["caption"])

    if report["sources"]:
        document.add_heading("Sources and References", level=1)
        for source in report["sources"]:
            document.add_paragraph(source)

    if report["warnings"]:
        document.add_heading("Limitations and Review Notes", level=1)
        for warning in report["warnings"]:
            document.add_paragraph(warning, style="List Bullet")

    document.save(output_path)

    return {
        "path": str(output_path.resolve()),
        "warnings": [],
    }
